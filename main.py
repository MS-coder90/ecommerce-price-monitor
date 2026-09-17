from scraper import get_page_html
from parser import parse_products
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

url = "https://webscraper.io/test-sites/e-commerce/static"

html = get_page_html(url)

products = parse_products(html)

print("Products found:", len(products))

df = pd.DataFrame(products)

df["scraped_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

df["price"] = df["price"].str.replace("$", "", regex=False).astype(float)

previous_file = "data/output/previous_prices.xlsx"

if __import__("os").path.exists(previous_file):
    old_df = pd.read_excel(previous_file)

    df = df.merge(
        old_df[["title", "price"]],
        on="title",
        how="left",
        suffixes=("", "_previous")
    )

    df["price_change"] = df["price"] - df["price_previous"]
    for _, row in df.iterrows():
     if pd.notna(row["price_change"]) and row["price_change"] != 0:
        if row["price_change"] > 0:
            print(
                f"PRICE INCREASE: {row['title']} "
                f"${row['price_previous']:.2f} -> ${row['price']:.2f} "
                f"(+${row['price_change']:.2f})"
            )
        else:
            print(
                f"PRICE DECREASE: {row['title']} "
                f"${row['price_previous']:.2f} -> ${row['price']:.2f} "
                f"(-${abs(row['price_change']):.2f})"
            )
else:
    df["price_previous"] = None
    df["price_change"] = None

df.to_excel(previous_file, index=False)

output_file = "data/output/live_products.xlsx"
df.to_excel(output_file, index=False)


# Format Excel report
wb = load_workbook(output_file)
ws = wb.active

# Make headers bold
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

# Format price columns
for row in range(2, ws.max_row + 1):
    ws[f"B{row}"].number_format = '$#,##0.00'
    ws[f"F{row}"].number_format = '$#,##0.00'

# Adjust column widths
for column in ws.columns:
    max_length = 0
    column_letter = get_column_letter(column[0].column)

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[column_letter].width = min(max_length + 2, 35)

# Freeze header
ws.freeze_panes = "A2"

# Add filter
ws.auto_filter.ref = ws.dimensions

wb.save(output_file)

print("\nExcel file created successfully!")
print(df)